"""Word 批量生成（邮件合并）：用 Excel 每一行填充 Word 模板里的 {{字段}}。

常见需求：批量生成合同、证书、通知书、工资条、录取通知。
用法:
  python -m office_toolkit.word_merge 模板.docx 数据.xlsx 输出文件夹 [--name-field 姓名]
模板里写 {{姓名}}、{{金额}} 等，字段名与 Excel 表头一致。
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

FIELD = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _fmt(value):
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if value.time() == datetime.min.time() else "%Y-%m-%d %H:%M")
    return str(value)


def _replace_in_paragraph(paragraph, row, missing):
    """Word 常把一个 {{字段}} 拆到多个 run 里，这里合并后替换，并保留第一个 run 的格式。"""
    full = "".join(r.text for r in paragraph.runs)
    if "{{" not in full:
        return

    def sub(m):
        key = m.group(1)
        if key not in row:
            missing.add(key)
            return m.group(0)
        return _fmt(row[key])

    new = FIELD.sub(sub, full)
    if new == full:
        return
    runs = paragraph.runs
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


def _iter_paragraphs(doc):
    def walk(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from walk(part)


def fill(template, row):
    doc = Document(str(template))
    missing = set()
    for p in _iter_paragraphs(doc):
        _replace_in_paragraph(p, row, missing)
    return doc, missing


def run(template, data, out_dir, name_field=None, sheet=0):
    df = pd.read_excel(data, sheet_name=sheet, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    all_missing = set()
    used = set()
    for i, rec in enumerate(df.to_dict("records"), 1):
        doc, missing = fill(template, rec)
        all_missing |= missing
        base = re.sub(r'[\\/:*?"<>|]', "_", _fmt(rec.get(name_field, ""))) if name_field else ""
        base = base or f"{i:04d}"
        name, k = base, 2
        while name in used:
            name = f"{base}_{k}"
            k += 1
        used.add(name)
        doc.save(out_dir / f"{name}.docx")
    return len(df), all_missing


def main(argv=None):
    ap = argparse.ArgumentParser(description="Word 模板批量生成")
    ap.add_argument("template"); ap.add_argument("data"); ap.add_argument("out_dir")
    ap.add_argument("--name-field", help="用哪一列做输出文件名")
    ap.add_argument("--sheet", default=0)
    a = ap.parse_args(argv)
    n, missing = run(a.template, a.data, a.out_dir, a.name_field, a.sheet)
    print(f"已生成 {n} 份文档 -> {a.out_dir}")
    if missing:
        print(f"注意：模板中这些字段在 Excel 表头里找不到，已原样保留: {sorted(missing)}")


if __name__ == "__main__":
    main()
