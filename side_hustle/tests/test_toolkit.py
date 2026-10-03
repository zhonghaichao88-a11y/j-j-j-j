import pandas as pd
import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter

from office_toolkit import excel_tools, image_tools, pdf_tools, rename_tools, word_merge


def test_excel_merge_and_split(tmp_path):
    src = tmp_path / "in"; src.mkdir()
    with pd.ExcelWriter(src / "a.xlsx") as w:
        pd.DataFrame({"部门": ["销售", "技术"], "金额": [1, 2]}).to_excel(w, sheet_name="一月", index=False)
        pd.DataFrame({"部门": ["销售"], "金额": [3]}).to_excel(w, sheet_name="二月", index=False)
    pd.DataFrame({"部门": ["财务/行政"], "金额": [4]}).to_csv(src / "b.csv", index=False, encoding="gbk")
    out = tmp_path / "merged.xlsx"
    assert excel_tools.merge(src, out) == (2, 4)
    df = pd.read_excel(out)
    assert list(df.columns[:2]) == ["来源文件", "来源工作表"]
    assert df["金额"].sum() == 10

    n = excel_tools.split(out, "部门", tmp_path / "split")
    assert n == 3
    assert (tmp_path / "split" / "财务_行政.xlsx").exists()
    excel_tools.split(out, "部门", tmp_path / "one", one_file=True)
    sheets = pd.read_excel(next((tmp_path / "one").iterdir()), sheet_name=None)
    assert len(sheets["销售"]) == 2


def _pdf(path, n):
    w = PdfWriter()
    for _ in range(n):
        w.add_blank_page(width=200, height=200)
    with open(path, "wb") as fh:
        w.write(fh)


def test_pdf_tools(tmp_path):
    _pdf(tmp_path / "a.pdf", 2); _pdf(tmp_path / "b.pdf", 3)
    out = tmp_path / "m.pdf"
    assert pdf_tools.merge([tmp_path], out) == 2
    assert len(PdfReader(out).pages) == 5
    assert pdf_tools.split(out, tmp_path / "s") == 5
    assert pdf_tools.parse_ranges("1-2,4，5-", 6) == [0, 1, 3, 4, 5]
    with pytest.raises(ValueError):
        pdf_tools.parse_ranges("7", 5)
    assert pdf_tools.pages(out, tmp_path / "p.pdf", "2-3") == 2


def test_word_merge_split_runs_tables_and_headers(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    for part in ["尊敬的 {{", "姓名", "}}，金额 {{金额}} 元，{{不存在}}"]:
        p.add_run(part)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "日期：{{日期}}"
    doc.sections[0].header.paragraphs[0].text = "编号 {{姓名}}"
    tpl = tmp_path / "t.docx"; doc.save(tpl)
    pd.DataFrame({"姓名": ["张三", "张三"], "金额": [100.0, 2.5],
                  "日期": pd.to_datetime(["2026-10-01", "2026-10-02"])}).to_excel(tmp_path / "d.xlsx", index=False)
    n, missing = word_merge.run(tpl, tmp_path / "d.xlsx", tmp_path / "out", name_field="姓名")
    assert n == 2 and missing == {"不存在"}
    d = Document(tmp_path / "out" / "张三.docx")
    assert d.paragraphs[0].text == "尊敬的 张三，金额 100 元，{{不存在}}"
    assert d.tables[0].cell(0, 0).text == "日期：2026-10-01"
    assert d.sections[0].header.paragraphs[0].text == "编号 张三"
    assert (tmp_path / "out" / "张三_2.docx").exists()


def test_image_batch(tmp_path):
    src = tmp_path / "src"; src.mkdir()
    Image.new("RGBA", (3000, 2000), (200, 30, 30, 255)).save(src / "x.png")
    n, before, after = image_tools.batch(src, tmp_path / "dst", max_side=1000, fmt="jpg",
                                         quality=70, watermark="版权所有")
    out = Image.open(tmp_path / "dst" / "x.jpg")
    assert n == 1 and out.size == (1000, 667) and out.mode == "RGB"


def test_rename_seq_and_map_with_swap(tmp_path):
    for name in ["b.txt", "a.txt"]:
        (tmp_path / name).write_text(name)
    plan, _ = rename_tools.plan_seq(tmp_path, "报告_", ext=".txt")
    assert [d.name for _, d in plan] == ["报告_001.txt", "报告_002.txt"]
    assert rename_tools.check(plan) == []
    rename_tools.apply(plan)
    assert (tmp_path / "报告_001.txt").read_text() == "a.txt"

    # 互换两个文件名
    pd.DataFrame({"原文件名": ["报告_001.txt", "报告_002.txt"],
                  "新文件名": ["报告_002", "报告_001"]}).to_excel(tmp_path / "map.xlsx", index=False)
    plan, problems = rename_tools.plan_map(tmp_path, tmp_path / "map.xlsx")
    assert problems == [] and rename_tools.check(plan) == []
    rename_tools.apply(plan)
    assert (tmp_path / "报告_001.txt").read_text() == "b.txt"


def test_rename_detects_conflicts(tmp_path):
    (tmp_path / "a.txt").write_text("a"); (tmp_path / "keep.txt").write_text("k")
    pd.DataFrame({"原文件名": ["a.txt", "缺失.txt"], "新文件名": ["keep.txt", "x"]}) \
        .to_excel(tmp_path / "m.xlsx", index=False)
    plan, problems = rename_tools.plan_map(tmp_path, tmp_path / "m.xlsx")
    assert problems == ["找不到: 缺失.txt"]
    assert rename_tools.check(plan) == ["目标已存在: keep.txt"]
